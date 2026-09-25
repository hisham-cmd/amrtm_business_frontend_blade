"""
Release Note Generator for AMRTM Business
Fills fields in the existing template only - no structural changes
"""

import os, sys
from datetime import datetime
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ARABIC_DAYS = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]

def set_cell_text(cell, text, bold=False, append=False):
    """Set text in a cell, preserving existing content if append=True"""
    if append:
        # Append to existing text (for title)
        for p in cell.paragraphs:
            for r in p.runs:
                r.text = r.text + text
        return
    
    # Clear cell paragraphs except first
    while len(cell.paragraphs) > 1:
        p = cell.paragraphs[-1]
        p._element.getparent().remove(p._element)
    
    p = cell.paragraphs[0]
    p.clear()
    
    run = p.add_run(text)
    run.bold = bold
    
    # Ensure RTL
    rPr = run._r.get_or_add_rPr()
    rPr.append(OxmlElement('w:rtl'))


def generate_release_note(template_path, output_path, version, date_str, time_str, day_str,
                          description, features, platform_type,
                          testers, developers, team_manager, upload_responsible, signatory,
                          version_title_prefix="إصدار رقم "):
    """Fill template fields only - no structural changes"""
    
    doc = Document(template_path)
    tables = doc.tables
    
    # P0 Title - append version to "إصدار رقم "
    for p in doc.paragraphs:
        if p.text.strip() == version_title_prefix.strip():
            for r in p.runs:
                r.text = version_title_prefix + version
            break
    
    # Table 0: نبذة عن الإصدار [0,1]
    if len(tables) > 0:
        set_cell_text(tables[0].cell(0, 1), description)
    
    # Table 1: Version info [0,1]=version [0,3]=day [1,1]=date [1,3]=time
    if len(tables) > 1:
        set_cell_text(tables[1].cell(0, 1), version, bold=True)
        set_cell_text(tables[1].cell(0, 3), day_str)
        set_cell_text(tables[1].cell(1, 1), date_str)
        set_cell_text(tables[1].cell(1, 3), time_str)
    
    # Table 2: Platform type [1,0]
    if len(tables) > 2:
        set_cell_text(tables[2].cell(1, 0), platform_type)
    
    # Table 3: Features [1,0]
    if len(tables) > 3:
        set_cell_text(tables[3].cell(1, 0), features)
    
    # Table 4: Testers & Developers - fill available rows only
    if len(tables) > 4:
        t = tables[4]
        max_data_rows = min(len(t.rows) - 1, max(len(testers), len(developers)))
        for i in range(max_data_rows):
            cell_1 = t.cell(i + 1, 1)
            cell_2 = t.cell(i + 1, 2)
            if i < len(testers):
                set_cell_text(cell_1, testers[i])
            if i < len(developers):
                set_cell_text(cell_2, developers[i])
    
    # Table 5: Team [0,1]=manager [1,1]=upload
    if len(tables) > 5:
        set_cell_text(tables[5].cell(0, 1), team_manager)
        set_cell_text(tables[5].cell(1, 1), upload_responsible)
    
    # Table 7: Signatory [0,1]
    if len(tables) > 7:
        set_cell_text(tables[7].cell(0, 1), signatory)
    
    # Save
    output_dir = os.path.dirname(output_path)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    doc.save(output_path)
    return output_path


def docx_to_pdf(docx_path, pdf_path):
    """Convert DOCX to PDF"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    
    font_name = 'Arial'
    try:
        pdfmetrics.registerFont(TTFont('Arial', 'C:/Windows/Fonts/arial.ttf'))
        pdfmetrics.registerFont(TTFont('Arial-Bold', 'C:/Windows/Fonts/arialbd.ttf'))
    except:
        font_name = 'Helvetica'

    read_doc = Document(docx_path)
    doc_builder = SimpleDocTemplate(pdf_path, pagesize=A4,
        rightMargin=2*cm, leftMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)

    story = []
    
    normal = ParagraphStyle('N', fontName=font_name, fontSize=11, leading=16, alignment=2, spaceAfter=6)
    heading = ParagraphStyle('H', fontName=font_name, fontSize=16, leading=24, alignment=1, spaceAfter=12)

    for p in read_doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style.name.startswith('Heading'):
            story.append(Paragraph(txt, heading))
        else:
            story.append(Paragraph(txt, normal))

    for table in read_doc.tables:
        data = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not data:
            continue
        col_count = len(data[0])
        avail = A4[0] - 4*cm
        tbl = Table(data, colWidths=[avail / col_count] * col_count)
        tbl.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 8))

    doc_builder.build(story)
    return pdf_path


def main():
    import argparse
    parser = argparse.ArgumentParser(description='Generate Release Note for AMRTM Business')
    parser.add_argument('--version', required=True)
    parser.add_argument('--date', required=True)
    parser.add_argument('--time', default=None)
    parser.add_argument('--description', default='')
    parser.add_argument('--features', default='')
    parser.add_argument('--platform', default='تطبيق ويب (Laravel)')
    parser.add_argument('--testers', default='')
    parser.add_argument('--developers', default='هشام الصغير')
    parser.add_argument('--team-manager', default='هشام قاسم الصغير')
    parser.add_argument('--upload-by', default='هشام الصغير')
    parser.add_argument('--signatory', default='هشام قاسم الصغير')
    parser.add_argument('--output', required=True)
    parser.add_argument('--pdf', default=None)

    args = parser.parse_args()

    # Template path for AMRTM Business
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    template_path = os.path.join(project_root, "إصدار رقم.docx")

    release_date = datetime.strptime(args.date, '%Y-%m-%d')
    day_str = ARABIC_DAYS[release_date.weekday()]
    time_str = args.time or release_date.strftime('%I:%M %p')

    testers = [t.strip() for t in args.testers.split(',') if t.strip()]
    developers = [d.strip() for d in args.developers.split(',') if d.strip()]

    print(f'Generating DOCX...')
    docx_path = generate_release_note(
        template_path, args.output, args.version, args.date, time_str, day_str,
        args.description, args.features, args.platform,
        testers, developers, args.team_manager, args.upload_by, args.signatory)
    print(f'DOCX: {docx_path}')

    if args.pdf:
        print(f'Generating PDF...')
        pdf_path = docx_to_pdf(docx_path, args.pdf)
        print(f'PDF: {pdf_path}')


if __name__ == '__main__':
    main()
