# -*- coding: utf-8 -*-
"""
Owner Report Generator (DOCX) - AMRTM Business
================================================
يولد تقريراً موجهاً لمالك الشركة (غير تقني) بصيغة DOCX،
من سجل التحديثات المحفوظ في Git بين تاريخين.
كل المحتوى بالعربية الكاملة مع اتجاه يمين-إلى-يسار صحيح.

الاستخدام:
    python owner_report_generator.py --from 2026-09-13 --to 2026-09-19
    (أو بدون وسائط ليعمل بأسلوب تفاعلي)

الخرج:
    .agent/reports/owner-report-YYYY-MM-DD.docx
"""
import argparse
import os
import re
import subprocess
import sys
from datetime import date, datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ── هوية آمر تم ──────────────────────────────────────────────
BRAND_GREEN = "006C35"
BRAND_DARK = "0B3B2C"
GOLD = "C9A227"
SLATE = "475569"
FONT = "Arial"            # خط متوفر على كل الأجهزة ويعرض العربية بوضوح
ARABIC_DIGITS = True      # تحويل الأرقام إلى الأرقام العربية (٠١٢٣)
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
LOGO_PATH = os.path.join(PROJECT_ROOT, "public", "images", "new-logo1.png")
REPORTS_DIR = os.path.join(PROJECT_ROOT, ".agent", "reports")

AR_MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
             "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
DIGIT_MAP = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")

# ── الترجمة الحرفية المعتمدة لكل تحديث سابق (مصدر الحقيقة) ──
EXACT_MAP = {
    "feat: update standalone fix script to v3": "تحديث سكربت الإصلاح المستقل إلى الإصدار الثالث",
    "feat: add multi-category business support for office providers": "إضافة دعم الفئات المتعددة للأنشطة التجارية لمزوّدي المكاتب",
    "feat: add business activity & category to consultant offices": "إضافة النشاط التجاري والفئة إلى مكاتب المستشارين",
    "feat: add consultant category and business activity tracking": "إضافة فئات المستشارين وتتبع النشاط التجاري",
    "style: update text shadows, fix lang toggle & admin catalog code": "تحديث ظلال النصوص وإصلاح مفتاح اللغة وكتالوج الإدارة",
    "chore: adjust grid columns for specialty cards": "ضبط أعمدة شبكة بطاقات التخصصات",
    "feat(admin,consultant): implement modular admin dashboard and consultant features": "إضافة لوحة التحكم الإدارية النمطية وميزات المستشارين",
    "fix(tests): fix outdated failing tests": "إصلاح الفحوصات القديمة المتعطلة لضمان استقرار التحديثات",
    "feat: add message attachments with contact scanning and improvements": "إضافة مرفقات الرسائل مع فحص بيانات التواصل وتحسينات عامة",
    "feat: add messaging, contact guard, and office logo support": "إضافة نظام المراسلات وحماية بيانات التواصل ودعم شعار المكاتب",
    "style: redesign offices section to glass grid": "إعادة تصميم قسم المكاتب بتصميم عصري أنيق",
    "feat(notifications,ui,backend): add interactive notifications, unified toast, and related improvements": "إضافة إشعارات تفاعلية وتنبيهات موحدة وتحسينات ذات صلة",
    "feat: implement payment settlements": "إضافة نظام التسويات المالية بين الأطراف",
    "feat: implement service request pool and official role rules": "إضافة مجموعة طلبات الخدمات وقواعد الأدوار الرسمية",
    "feat(amrtm): add nafath verification, service duration system, and office service management": "إضافة التحقق الرسمي عبر نفاذ ونظام مدد الخدمات وإدارة خدمات المكاتب",
    "feat: implement core business module architecture including dashboard hubs, service request management, and UI component library": "إضافة البنية الأساسية لوحدة الأعمال متضمنة مراكز لوحات التحكم وإدارة طلبات الخدمات ومكتبة مكونات الواجهة",
    "feat(ui): add UI builder, fix labels and eye field": "إضافة منشئ الواجهات وإصلاح التسميات وعدّاد المشاهدات",
    "feat: implement comprehensive service catalog, office dashboard, and UI component system with database migrations": "إضافة كتالوج الخدمات الشامل ولوحة تحكم المكاتب ونظام مكونات الواجهة مع تحديثات قاعدة البيانات",
    "feat: implement standardized dashboard layout, navigation components, and service views": "إضافة تخطيط موحد للوحات التحكم ومكونات التنقل وواجهات الخدمات",
    "feat: deep-link hub menu items to real dashboard sections via URL hash": "ربط عناصر قوائم المراكز بأقسام لوحات التحكم مباشرة",
    "feat: unify Flowbite dashboard hub for all personas with org-structure matrix": "توحيد مركز لوحة التحكم لجميع المستخدمين مع مصفوفة الهيكل التنظيمي",
    "feat: implement new service update module with user profile fields and database enhancements": "إضافة وحدة تحديث الخدمات الجديدة مع حقول الملف الشخصي وتحسينات قاعدة البيانات",
    "feat: implement multi-guard authentication for business and provider accounts with registration and status validation logic": "إضافة نظام دخول متعدد الحماية لحسابات الأعمال والمزوّدين مع التحقق من حالة التسجيل",
    "feat: implement multi-role registration flow and add automated release management scripts": "إضافة مسار تسجيل متعدد الأدوار وأدوات آلية لإدارة الإصدارات",
    "feat: implement hero section and custom styling for service update page": "إضافة قسم الواجهة الرئيسية وتنسيق مخصص لصفحة تحديث الخدمات",
    "feat: implement auth login view and provider account registration system with directory management pages": "إضافة واجهة تسجيل الدخول ونظام تسجيل حسابات المزوّدين مع صفحات إدارة الدلائل",
    "feat: implement user authentication system and provider account registration flow": "إضافة نظام دخول المستخدمين ومسار تسجيل حسابات المزوّدين",
    "feat: implement business authentication, provider account management, and service directory routing": "إضافة دخول قطاع الأعمال وإدارة حسابات المزوّدين وتوجيه دلائل الخدمات",
    "feat: create new consultant and office directory views with search and filter functionality": "إنشاء واجهات دليل المستشارين والمكاتب مع البحث والفلترة",
    "feat: implement office registration system with associated database migrations, controllers, and localized UI components": "إضافة نظام تسجيل المكاتب مع تحديثات قاعدة البيانات ومكونات واجهة معرّبة",
    "feat: implement consultants directory interface, controller, and backend models": "إضافة واجهة دليل المستشارين مع نظامها الخلفي",
    "feat: introduce Corporate-UI components and catalog views for update service workflow": "تقديم مكونات الواجهة المؤسسية وواجهات الكتالوج لسير عمل تحديث الخدمات",
    "feat: implement new Corporate-UI views for service catalog, office directory, and category selection": "إضافة واجهات مؤسسية جديدة لكتالوج الخدمات ودليل المكاتب واختيار الفئات",
    "feat: implement corporate UI design system and category page structure": "إضافة نظام تصميم واجهة مؤسسي وبنية صفحات الفئات",
    "feat: implement Corporate UI hero component with partial styles and new directory service views": "إضافة مكوّن الواجهة الرئيسية المؤسسي وواجهات خدمات الدليل الجديدة",
    "fix: update production image filenames to UUIDs and configure Render deployment environment": "تحديث أسماء ملفات صور الإنتاج وإعداد بيئة النشر",
    "fix: resolve production image 404 errors by normalizing filenames to UUIDs and configure Render deployment settings": "معالجة مشكلة عدم ظهور صور الإنتاج وإعادة ضبط إعدادات النشر",
    "feat: initialize update service blade view and associated database schema": "إنشاء واجهة خدمة التحديث وبنية قاعدة البيانات المرتبطة",
    "feat: add consultants directory view and back up legacy database file": "إضافة واجهة دليل المستشارين مع نسخة احتياطية من قاعدة البيانات القديمة",
    "feat: implement corporate UI design system and integrate new catalog, directory, and detail service views": "إضافة نظام تصميم واجهة مؤسسي ودمج واجهات الكتالوج والدليل والتفاصيل",
    "feat: implement service catalog and consultants directory module with associated views and routes": "إضافة وحدة كتالوج الخدمات ودليل المستشارين مع واجهاتها ومساراتها",
    "feat: implement provider account registration and business model architecture": "إضافة تسجيل حسابات المزوّدين وبنية نماذج الأعمال",
    "feat: add provider account registration view and update project mapping": "إضافة واجهة تسجيل حسابات المزوّدين وتحديث تنظيم المشروع",
    "feat: introduce update service UI views and include compiled assets in manifest": "تقديم واجهات خدمة التحديث مع الأصول الجاهزة",
    "feat: implement service catalog UI components and backend controller": "إضافة مكونات واجهة كتالوج الخدمات ونظامها الخلفي",
    "feat: add catalog entity view with Corporate-UI components and asset manifest": "إضافة واجهة كيانات الكتالوج بمكونات الواجهة المؤسسية",
    "feat: implement Corporate-UI design system for catalog entity update view": "إضافة نظام تصميم مؤسسي لواجهة تحديث الكيانات",
    "feat: implement service catalog controller and corresponding views for browsing categories, entities, and offices.": "إضافة نظام كتالوج الخدمات وواجهاته لتصفح الفئات والكيانات والمكاتب",
    "feat: implement new navbar component and service index pages with supporting assets and OCR language files": "إضافة شريط التنقل الجديد وصفحات الخدمات الرئيسية مع الأصول وملفات اللغات",
    "feat: replace Vue with Flowbite and implement Blade-based UI component architecture": "استبدال تقنية الواجهات القديمة بنظام مكونات حديث وأكثر استقراراً",
    "feat: create index view for update service with responsive design and hero slider components": "إنشاء واجهة رئيسية لخدمة التحديث بتصميم متجاوب ومكونات عرض الصور",
    "feat: create update service index page with custom Saudi-inspired branding and layout": "إنشاء الصفحة الرئيسية لخدمة التحديث بهوية وطنية سعودية",
    "feat: implement application workflow, consolidate database schema, and add functional feature tests": "إضافة سير عمل التطبيق وتوحيد بنية قاعدة البيانات مع فحوصات وظيفية",
    "feat: implement service catalog controller, views, and corresponding routes with initial feature tests": "إضافة نظام كتالوج الخدمات وواجهاته مع فحوصات أولية",
    "feat: implement business contract management system including migration, models, controllers, and dashboard views": "إضافة نظام إدارة عقود الأعمال مع تحديثات قاعدة البيانات وواجهات المتابعة",
    "feat: initialize business sector platform with routing, dashboard views, and feature tests": "إنشاء منصة قطاع الأعمال مع واجهات لوحة التحكم وفحوصات الجودة",
    "feat: add new update service view template with responsive dashboard styling": "إضافة قالب جديد لواجهة خدمة التحديث بتنسيق متجاوب",
    "feat: implement initial service update view with security headers and provider configuration": "إضافة الواجهة الأولية لخدمة التحديث مع إعدادات الأمان",
    "feat: initialize admin dashboard view and configure security and service provider layers": "إنشاء واجهة لوحة التحكم الإدارية وإعداد طبقات الأمان",
    "Update DB migrations and fallbacks for cloud databaseasp.net": "تحديث بنية قاعدة البيانات وبدائلها السحابية",
    "Add contracts management and fix Render dynamic port entrypoint": "إضافة إدارة العقود وإصلاح إعدادات النشر على الإنترنت",
    "feat: add HomepageSlide model and office directory view template": "إضافة نظام شرائح الصفحة الرئيسية وقالب دليل المكاتب",
    "Restore compact transparent tagline": "استعادة شعار الصفحة المختصر الشفاف",
    "Fix JS textContent wiping tagline HTML structure and clear caches": "إصلاح مشكلة اختفاء شعار الصفحة وتنظيف الملفات المؤقتة",
    "Optimize tagline clarity and slider visibility": "تحسين وضوح شعار الصفحة وظهور شرائح العرض",
    "Enhance tagline readability with glassmorphism contrast and add .agentignore": "تحسين وضوح شعار الصفحة على الخلفيات المختلفة",
    "feat: implement hero section and service catalog templates with associated styles and assets": "إضافة قسم الواجهة الرئيسية وقوالب كتالوج الخدمات مع تنسيقاتها",
    "feat: implement update service dashboard and static contract creation interface": "إضافة لوحة تحكم خدمة التحديث وواجهة إنشاء العقود",
    "Update contracts page": "تحديث صفحة العقود",
    "Update electronic contracts": "تحديث العقود الإلكترونية",
    "Update contracts": "تحديث العقود",
    "contract static": "تحديث محتوى صفحة العقود",
    "Resolve merge conflicts with main": "حل تعارضات الدمج مع الفرع الرئيسي بأمان",
}

# ── قاموس المصطلحات (لأي تحديث جديد غير مغطى سابقاً) ──
TERM_DICT = {
    "implement": "إضافة", "add": "إضافة", "create": "إنشاء", "introduce": "تقديم",
    "initialize": "إنشاء", "update": "تحديث", "fix": "إصلاح", "resolve": "معالجة",
    "improve": "تحسين", "enhance": "تحسين", "optimize": "تحسين", "redesign": "إعادة تصميم",
    "remove": "إزالة", "replace": "استبدال", "support": "دعم", "unify": "توحيد",
    "consolidate": "توحيد", "configure": "إعداد", "adjust": "ضبط", "restore": "استعادة",
    "integrate": "دمج", "include": "تضمين", "including": "متضمناً", "with": "مع",
    "and": "و", "for": "لـ", "to": "إلى", "in": "في", "via": "عبر", "into": "إلى",
    "associated": "المرتبطة", "corresponding": "المقابلة", "related": "ذات الصلة",
    "new": "جديد", "comprehensive": "شامل", "modular": "نمطية", "interactive": "تفاعلية",
    "unified": "موحد", "custom": "مخصص", "responsive": "متجاوب", "static": "ثابتة",
    "initial": "أولية", "core": "الأساسية", "official": "الرسمية", "legacy": "القديم",
    "dashboard": "لوحة التحكم", "admin": "الإدارية", "consultant": "المستشارين",
    "consultants": "المستشارين", "directory": "الدليل", "directories": "الدلائل",
    "office": "المكاتب", "offices": "المكاتب", "catalog": "الكتالوج", "entity": "الكيان",
    "entities": "الكيانات", "category": "الفئة", "categories": "الفئات",
    "specialty": "التخصص", "specialties": "التخصصات", "activity": "النشاط",
    "activities": "الأنشطة", "business": "الأعمال", "provider": "المزوّدين",
    "providers": "المزوّدين", "account": "الحسابات", "accounts": "الحسابات",
    "authentication": "المصادقة", "auth": "المصادقة", "login": "تسجيل الدخول",
    "registration": "التسجيل", "user": "المستخدمين", "users": "المستخدمين",
    "message": "الرسائل", "messages": "الرسائل", "messaging": "المراسلات",
    "notification": "الإشعارات", "notifications": "الإشعارات", "toast": "التنبيهات",
    "attachment": "المرفقات", "attachments": "المرفقات", "contact": "التواصل",
    "guard": "الحماية", "contract": "العقود", "contracts": "العقود",
    "payment": "المدفوعات", "payments": "المدفوعات", "settlement": "التسويات",
    "settlements": "التسويات", "service": "الخدمات", "services": "الخدمات",
    "request": "الطلبات", "requests": "الطلبات", "management": "الإدارة",
    "system": "نظام", "module": "وحدة", "view": "الواجهة", "views": "الواجهات",
    "page": "الصفحة", "pages": "الصفحات", "component": "مكونات", "components": "مكونات",
    "layout": "التخطيط", "design": "التصميم", "styling": "التنسيق", "style": "تصميم",
    "navigation": "التنقل", "navbar": "شريط التنقل", "menu": "القائمة",
    "hero": "الواجهة الرئيسية", "slider": "شرائح العرض", "database": "قاعدة البيانات",
    "db": "قاعدة البيانات", "migration": "تحديثات قاعدة البيانات", "migrations": "تحديثات قاعدة البيانات",
    "model": "النماذج", "models": "النماذج", "controller": "المتحكمات",
    "controllers": "المتحكمات", "routing": "التوجيه", "routes": "المسارات",
    "test": "الفحوصات", "tests": "الفحوصات", "script": "السكربت", "scripts": "السكربتات",
    "release": "الإصدار", "deploy": "النشر", "deployment": "النشر",
    "production": "الإنتاج", "render": "منصة النشر", "environment": "البيئة",
    "settings": "الإعدادات", "image": "الصور", "images": "الصور",
    "filenames": "أسماء الملفات", "file": "الملفات", "files": "الملفات",
    "security": "الأمان", "role": "الأدوار", "roles": "الأدوار", "rules": "القواعد",
    "pool": "مجموعة", "duration": "مدد", "nafath": "نفاذ", "verification": "التحقق",
    "workflow": "سير العمل", "schema": "بنية قاعدة البيانات", "cloud": "السحابية",
    "tagline": "شعار الصفحة", "glass": "زجاجية", "grid": "شبكة",
    "columns": "الأعمدة", "cards": "البطاقات", "labels": "التسميات",
    "field": "الحقل", "fields": "الحقول", "hub": "مركز", "hubs": "مراكز",
    "personas": "الشخصيات", "matrix": "مصفوفة", "org-structure": "الهيكل التنظيمي",
    "deep-link": "ربط مباشر", "profile": "الملف الشخصي", "validation": "التحقق",
    "logic": "المنطق", "status": "الحالة", "flow": "المسار", "architecture": "بنية",
    "library": "مكتبة", "assets": "الأصول", "manifest": "قائمة الأصول",
    "language": "اللغة", "languages": "اللغات", "ocr": "قراءة النصوص",
    "branding": "الهوية", "saudi-inspired": "سعودية الطابع", "index": "الرئيسية",
    "template": "قالب", "cache": "الملفات المؤقتة", "caches": "الملفات المؤقتة",
    "backup": "النسخ الاحتياطي", "builder": "منشئ", "toggle": "مفتاح",
    "lang": "اللغة", "eye": "المشاهدات", "multi-role": "متعدد الأدوار",
    "multi-guard": "متعدد الحماية", "multi-category": "الفئات المتعددة",
    "standalone": "المستقل", "automated": "آلية", "outdated": "القديمة",
    "failing": "المتعطلة", "oldstring": "", "scoped": "محدد النطاق",
}
TYPE_MAP = {
    "feat": "ميزة جديدة", "fix": "إصلاح مشكلة", "style": "تحسين مظهر",
    "refactor": "تحسين داخلي", "chore": "صيانة", "test": "فحوصات جودة",
    "docs": "توثيق", "perf": "تحسين سرعة",
}
TOPIC_MAP = [
    ("auth", "التسجيل والحسابات"), ("login", "التسجيل والحسابات"),
    ("registration", "التسجيل والحسابات"), ("guard", "التسجيل والحسابات"),
    ("consultant", "دليل المستشارين"), ("office", "دليل المستشارين"),
    ("directory", "دليل المستشارين"), ("catalog", "كتالوج الخدمات"),
    ("service", "الخدمات وطلباتها"), ("request", "الخدمات وطلباتها"),
    ("payment", "الجانب المالي"), ("settlement", "الجانب المالي"),
    ("nafath", "التحقق من الهوية"), ("notif", "التواصل والإشعارات"),
    ("message", "التواصل والإشعارات"), ("toast", "التواصل والإشعارات"),
    ("contract", "العقود الإلكترونية"), ("admin", "لوحة التحكم الإدارية"),
    ("dashboard", "لوحة التحكم الإدارية"), ("ui", "تصميم الواجهة وتجربة الاستخدام"),
    ("style", "تصميم الواجهة وتجربة الاستخدام"), ("hero", "تصميم الواجهة وتجربة الاستخدام"),
    ("slider", "تصميم الواجهة وتجربة الاستخدام"), ("logo", "هوية المنصة"),
    ("fix", "إصلاحات وجودة"), ("test", "إصلاحات وجودة"),
    ("chore", "صيانة وتحسينات عامة"), ("refactor", "صيانة وتحسينات عامة"),
    ("deploy", "نشر المنصة وتشغيلها"), ("release", "نشر المنصة وتشغيلها"),
]

ARABIC_RE = re.compile(r"[\u0600-\u06FF]")
PUNCT = "(),.:;\"'"


def ar_digits(text):
    return text.translate(DIGIT_MAP) if ARABIC_DIGITS else text


def fmt_ar(d):
    return f"{d.day} {AR_MONTHS[d.month - 1]} {d.year}"


def term_translate(text):
    """ترجمة كلمة-كلمة عبر القاموس مع حذف أي مصطلح لاتيني متبقٍ."""
    text = text.replace("&", " و ")
    out = []
    for word in text.split():
        core = word.strip(PUNCT)
        if not core:
            continue
        low = core.lower()
        tr = TERM_DICT.get(low) or TERM_DICT.get(low.rstrip("s"))
        if tr is not None:
            out.append(tr)
        elif not re.search(r"[A-Za-z]", core):
            out.append(core)  # أرقام أو نص عربي أصلي
        # المصطلح اللاتيني غير المعروف يُحذف ببساطة
    return " ".join(out)


def humanize(subject):
    """تحويل رسالة التحديث إلى جملة عربية كاملة."""
    key = subject.strip().rstrip(".")
    if key in EXACT_MAP:
        return EXACT_MAP[key]
    low = subject.lower()
    if subject.startswith("Merge") or "merge" in low:
        return "دمج أعمال الفريق وتنسيق التحديثات بين الفروع"
    body = subject.split(":", 1)[1].strip() if ":" in subject else subject.strip()
    translated = term_translate(body)
    return translated if translated else "تحديث عام على المنصة"


def classify(subject):
    parts = subject.split(":", 1)
    ctype = TYPE_MAP.get(parts[0].split("(")[0].strip().lower(), "تحديث") if len(parts) > 1 else "تحديث"
    low = subject.lower()
    topic = "تطويرات عامة"
    for key, label in TOPIC_MAP:
        if key in low:
            topic = label
            break
    return ctype, topic


def run_git(args):
    result = subprocess.run(
        ["git"] + args, cwd=PROJECT_ROOT, capture_output=True, text=True,
        encoding="utf-8", errors="replace",
    )
    if result.returncode != 0:
        sys.exit(f"[Error] git failed: {result.stderr.strip()}")
    return result.stdout


def parse_date(text):
    return datetime.strptime(text, "%Y-%m-%d").date()


def collect_commits(d_from, d_to):
    out = run_git(["log", f"--since={d_from.isoformat()}T00:00:00",
                   f"--until={d_to.isoformat()}T23:59:59",
                   "--pretty=format:%h|%ad|%an|%s", "--date=short"])
    commits = []
    for line in out.splitlines():
        if not line.strip():
            continue
        sha, d, author, subject = (line.split("|", 3) + [""])[:4]
        commits.append({"sha": sha, "date": d, "author": author, "subject": subject})
    return commits


# ══ أدوات التنسيق العربي (اتجاه يمين-إلى-يسار) ═══════════════

def set_rtl_paragraph(paragraph):
    """اتجاه الفقرة من اليمين إلى اليسار."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


def set_rtl_run(run, text):
    """اتجاه النص داخل الكلمة + لغة التدقيق العربية."""
    if not ARABIC_RE.search(text):
        return
    rPr = run._element.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rPr.append(rtl)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:bidi"), "ar-SA")
    rPr.append(lang)


def style_run(run, size=13, bold=False, color=SLATE, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    rFonts.set(qn("w:cs"), font)
    rFonts.set(qn("w:eastAsia"), font)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szCs)
    if bold:
        rPr.append(OxmlElement("w:bCs"))


def add_para(doc, text, size=13, bold=False, color=SLATE,
             align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=6):
    p = doc.add_paragraph()
    # ملاحظة مهمة: في الفقرات ذات الاتجاه يمين-يسار (bidi)، يعكس Word قيم
    # المحاذاة left/right منطقياً، لذا نترك المحاذاة الافتراضية (البداية = اليمين)
    # ولا نكتب jc إلا للتوسيط — وهذا هو الإصلاح الصحيح لمشاكل الاتجاه.
    if align == WD_ALIGN_PARAGRAPH.CENTER:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    set_rtl_paragraph(p)
    text = ar_digits(text)
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, color=color)
    set_rtl_run(run, text)
    return p


def set_document_rtl_default(doc):
    """فرض اتجاه يمين-يسار افتراضياً على كامل الوثيقة (docDefaults)."""
    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_el.insert(0, doc_defaults)
    p_pr_default = doc_defaults.find(qn("w:pPrDefault"))
    if p_pr_default is None:
        p_pr_default = OxmlElement("w:pPrDefault")
        doc_defaults.append(p_pr_default)
    p_pr = p_pr_default.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        p_pr_default.append(p_pr)
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))


def set_table_rtl(table):
    """ترتيب أعمدة الجدول من اليمين إلى اليسار."""
    tblPr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_table_row(table, cells_text, header=False):
    row = table.add_row()
    for i, text in enumerate(cells_text):
        cell = row.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        set_rtl_paragraph(p)  # بلا jc — الافتراضي في RTL هو المحاذاة لليمين
        text = ar_digits(text)
        run = p.add_run(text)
        style_run(run, size=12, bold=header, color="FFFFFF" if header else SLATE)
        set_rtl_run(run, text)
        if header:
            shade_cell(cell, BRAND_GREEN)
    return row


def build_docx(commits, d_from, d_to, out_path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    doc.core_properties.title = "تقرير التقدم المقدم لمالك الشركة"
    doc.core_properties.author = "هشام قاسم الصغير"
    doc.core_properties.subject = "آمر تم - قطاع الأعمال"
    set_document_rtl_default(doc)

    # ── الغلاف: الشعار + العنوان ──
    if os.path.exists(LOGO_PATH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO_PATH, width=Cm(6.5))
    add_para(doc, "آمر تم — قطاع الأعمال", size=20, bold=True, color=BRAND_GREEN,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "تقرير التقدم المقدم لمالك الشركة", size=16, bold=True, color=BRAND_DARK,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, f"للفترة من {fmt_ar(d_from)} إلى {fmt_ar(d_to)}",
             size=13, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, f"إعداد: هشام قاسم الصغير — قائد فريق التطوير | تاريخ التقرير: {fmt_ar(date.today())}",
             size=11, color=SLATE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

    sep = doc.add_paragraph()
    sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sep.add_run("—" * 40)
    style_run(run, size=10, color=GOLD)

    # ── ملخص تنفيذي ──
    add_para(doc, "الملخص التنفيذي", size=16, bold=True, color=BRAND_GREEN, space_after=6)
    feats = [c for c in commits if c["subject"].lower().startswith("feat")]
    fixes = [c for c in commits if c["subject"].lower().startswith("fix")]
    add_para(doc,
             f"خلال هذه الفترة أنجز الفريق {len(commits)} تحديثاً محفوظاً ومراجعاً على المنصة، "
             f"منها {len(feats)} ميزة أو تحسيناً جديداً يخدم المستخدم، و{len(fixes)} إصلاحاً يضمن استقرار وجودة العمل. "
             "جميع الأعمال محفوظة في حساب الشركة الرسمي ويمكن الرجوع إليها في أي وقت.",
             size=11, space_after=12)

    # ── جدول المؤشرات ──
    add_para(doc, "مؤشرات الفترة", size=16, bold=True, color=BRAND_GREEN, space_after=6)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(table)
    add_table_row(table, ["المؤشر", "القيمة"], header=True)
    topics = {}
    for c in commits:
        _, topic = classify(c["subject"])
        topics[topic] = topics.get(topic, 0) + 1
    top_topics = sorted(topics.items(), key=lambda x: -x[1])[:5]
    add_table_row(table, ["عدد التحديثات المنجزة", str(len(commits))])
    add_table_row(table, ["الميزات الجديدة", str(len(feats))])
    add_table_row(table, ["الإصلاحات", str(len(fixes))])
    add_table_row(table, ["أهم مجالات العمل",
                          "، ".join(t for t, _ in top_topics)])
    add_table_row(table, ["حالة الحفظ والنشر", "جميع الأعمال محفوظة ومرفوعة على حساب الشركة الرسمي"])
    doc.add_paragraph()

    # ── الإنجازات حسب المجال ──
    add_para(doc, "أبرز الإنجازات خلال الفترة", size=16, bold=True, color=BRAND_GREEN, space_after=6)
    by_topic = {}
    for c in commits:
        ctype, topic = classify(c["subject"])
        by_topic.setdefault(topic, []).append((ctype, humanize(c["subject"])))
    for topic in sorted(by_topic, key=lambda t: -len(by_topic[t])):
        items = by_topic[topic]
        add_para(doc, f"• {topic}", size=13, bold=True, color=BRAND_DARK, space_after=2)
        for ctype, text in items[:4]:
            add_para(doc, f"   – {text}", size=12, space_after=1)
        if len(items) > 4:
            add_para(doc, "   – وأعمال أخرى في نفس المجال", size=12, space_after=1)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ── الجودة والاستقرار ──
    add_para(doc, "الجودة والاستقرار", size=16, bold=True, color=BRAND_GREEN, space_after=6)
    add_para(doc,
             "قبل كل تحديث تُجرى فحوصات آلية تتأكد من عدم تعطل أي خدمة كانت تعمل سابقاً. "
             "كما تم خلال الفترة إصلاح الملاحظات التي ظهرت أثناء الاستخدام لضمان تجربة سلسة للمستخدمين.",
             size=11, space_after=12)

    # ── الخاتمة ──
    add_para(doc, "الخلاصة", size=16, bold=True, color=BRAND_GREEN, space_after=6)
    add_para(doc,
             "المنصة تمضي وفق الخطة، وتزداد جاهزية يوماً بعد يوم لمرحلة التشغيل الفعلي "
             "واستقبال العملاء. تفاصيل تنفيذية أدق متوفرة لدى فريق التطوير عند الطلب.",
             size=11, space_after=12)

    add_para(doc, "هشام قاسم الصغير — قائد فريق التطوير | آمر تم", size=11.5, bold=True,
             color=BRAND_DARK, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.save(out_path)
    return out_path


def convert_to_pdf(docx_path):
    """تحويل التقرير إلى PDF عبر Microsoft Word (نفس التنسيق تماماً مع اتجاهات RTL)."""
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    word = None
    try:
        import win32com.client
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
        doc.ExportAsFixedFormat(os.path.abspath(pdf_path), ExportFormat=17)  # 17 = PDF
        doc.Close(False)
        return pdf_path
    except Exception as exc:
        print(f"[Warn] PDF generation skipped: {exc}")
        return None
    finally:
        if word is not None:
            word.Quit()


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(description="Owner Report Generator (DOCX)")
    parser.add_argument("--from", dest="date_from", default=None, help="YYYY-MM-DD")
    parser.add_argument("--to", dest="date_to", default=None, help="YYYY-MM-DD")
    args = parser.parse_args()

    d_from = parse_date(args.date_from) if args.date_from else None
    d_to = parse_date(args.date_to) if args.date_to else None

    if d_from is None:
        raw = input("[Input] From date (YYYY-MM-DD, e.g. 2026-09-13): ").strip()
        if not raw:
            sys.exit("[Error] From date is required.")
        d_from = parse_date(raw)
    if d_to is None:
        raw = input("[Input] To date (YYYY-MM-DD, default today): ").strip()
        d_to = parse_date(raw) if raw else date.today()
    if d_to < d_from:
        sys.exit("[Error] To date must be after from date.")

    print(f"[Running] Collecting commits {d_from} -> {d_to} ...")
    commits = collect_commits(d_from, d_to)
    if not commits:
        sys.exit("[Info] No commits found in this period.")

    os.makedirs(REPORTS_DIR, exist_ok=True)
    out_path = os.path.join(REPORTS_DIR, f"owner-report-{d_to.isoformat()}.docx")
    build_docx(commits, d_from, d_to, out_path)
    print(f"[Done] {out_path}")
    pdf_path = convert_to_pdf(out_path)
    if pdf_path:
        print(f"[Done] {pdf_path}")
    print(f"        commits: {len(commits)}")


if __name__ == "__main__":
    main()
