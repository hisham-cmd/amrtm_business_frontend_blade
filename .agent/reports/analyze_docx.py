from docx import Document
import os

# AMRTM Business project
project_root = r'C:\react_projects\amrtm_business'
docx_path = os.path.join(project_root, '.agent', 'reports', 'إصدار رقم 1.0.0.docx')

doc = Document(docx_path)
print('=== PARAGRAPHS ===')
for i, p in enumerate(doc.paragraphs):
    print(f'{i}: [{p.style.name}] {p.text[:100]}')
print()
print('=== TABLES ===')
for ti, table in enumerate(doc.tables):
    print(f'Table {ti}: {len(table.rows)} rows x {len(table.columns)} cols')
    for ri, row in enumerate(table.rows):
        cells = [cell.text[:30] for cell in row.cells]
        print(f'  Row {ri}: {cells}')
