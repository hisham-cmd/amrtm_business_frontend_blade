import sys, json, os
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

project_root = r'C:\react_projects\amrtm_business'
doc = Document(os.path.join(project_root, 'إصدار رقم.docx'))
result = {'paragraphs': [], 'tables': []}
for i, p in enumerate(doc.paragraphs):
    result['paragraphs'].append({'i': i, 'style': p.style.name, 'text': p.text[:100]})
for ti, table in enumerate(doc.tables):
    t = {'index': ti, 'rows': len(table.rows), 'cols': len(table.columns), 'data': []}
    for ri, row in enumerate(table.rows):
        cells = [cell.text[:40] for cell in row.cells]
        t['data'].append(cells)
    result['tables'].append(t)
print(json.dumps(result, ensure_ascii=False, indent=2))
